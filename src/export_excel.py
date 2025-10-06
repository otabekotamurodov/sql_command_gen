import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from datetime import datetime
from run_query import run_query
from llm_query_generator import generate_sql
import os


def export_to_excel(df: pd.DataFrame, filename: str = None):
    if df is None or df.empty:
        raise ValueError("DataFrame bo'sh. Export qilish uchun ma'lumot yo'q.")

    os.makedirs("results", exist_ok=True)

    if filename is None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"results/query_result_{timestamp}.xlsx"

    # write DataFrame
    df.to_excel(filename, index=False, sheet_name="Results")

    # add Diagramms
    wb = load_workbook(filename)
    ws = wb["Results"]

    # Agar faqat bitta ustun bulsa grafik yaratilmaydi
    if df.shape[1] >= 2:
        chart = BarChart()
        chart.title = "Query Result Chart"
        chart.x_axis.title = df.columns[0]
        chart.y_axis.title = "Qiymat"

        # Ma'lumot diapazoni (bitta kategoriyali yoki bir nechta)
        data_ref = Reference(ws, min_col=2, min_row=1, max_col=df.shape[1], max_row=df.shape[0] + 1)
        cats_ref = Reference(ws, min_col=1, min_row=2, max_row=df.shape[0] + 1)

        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)

        ws.add_chart(chart, f"E2")  # Diagrammani joylash joyi
    else:
        print("⚠️ Diagramma uchun ustunlar yetarli emas, faqat jadval yozildi.")

    wb.save(filename)
    wb.close()

    print(f"✅ Excel fayl saqlandi: {filename}")
    return filename


# ---- CLI test ----
if __name__ == "__main__":
    prompt = "2024 yilida oylar kesimida Toshkent viloyati bo‘yicha jami tranzaksiyalar summasini ko‘rsat"
    print("🧠 Prompt:", prompt)

    sql = generate_sql(prompt)
    df = run_query(sql)
    export_to_excel(df)
